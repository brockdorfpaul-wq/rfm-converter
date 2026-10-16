#!/usr/bin/env python3
"""
Convert a wide "Item Prices" workbook (one row per restaurant, an In-Store/Delivery
column pair per item) into the long "RFM Data" format:

    Store ID | Item ID | In-Store Recommended Price | Delivery Recommended Price

Usage:
    python convert_to_rfm.py McRib_16-October-2026.xlsx
    python convert_to_rfm.py input.xlsx -o output_RFM_Format.xlsx
"""
import argparse
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

HEADERS = ["Store ID", "Item ID", "In-Store Recommended Price", "Delivery Recommended Price"]


def find_header_rows(ws):
    """Locate the item-group row (cells like '#10\\nMcRib') and the column-header row (NSN)."""
    item_row = nsn_row = None
    for r in range(1, ws.max_row + 1):
        values = [c.value for c in ws[r]]
        if nsn_row is None and "NSN" in values:
            nsn_row = r
        if item_row is None and any(isinstance(v, str) and re.match(r"\s*#\s*\d+", v) for v in values):
            item_row = r
    if item_row is None or nsn_row is None:
        raise ValueError("Could not find the item-ID row or the NSN header row.")
    return item_row, nsn_row


def read_wide(path, sheet=None):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    item_row, nsn_row = find_header_rows(ws)

    nsn_col = next(c.column for c in ws[nsn_row] if c.value == "NSN")

    # Map each item to its In-Store / Delivery columns using the sub-headers.
    items = []  # (item_id, instore_col, delivery_col)
    for cell in ws[item_row]:
        if isinstance(cell.value, str):
            m = re.match(r"\s*#\s*(\d+)", cell.value)
            if not m:
                continue
            item_id, start = m.group(1), cell.column
            in_col = del_col = None
            for col in (start, start + 1):
                label = str(ws.cell(nsn_row, col).value or "").strip().lower()
                if label.startswith("in-store"):
                    in_col = col
                elif label.startswith("delivery"):
                    del_col = col
            if in_col is None and del_col is None:
                raise ValueError(f"Item #{item_id}: no In-Store/Delivery columns found under it.")
            if in_col is None or del_col is None:
                print(f"Warning: item #{item_id} has only one price column; the other will be blank.")
            items.append((item_id, in_col, del_col))

    rows = []
    for r in range(nsn_row + 1, ws.max_row + 1):
        nsn = ws.cell(r, nsn_col).value
        if nsn in (None, ""):
            continue
        store_id = int(nsn)
        for item_id, in_col, del_col in items:
            rows.append((store_id, item_id,
                         ws.cell(r, in_col).value if in_col else None,
                         ws.cell(r, del_col).value if del_col else None))

    # Target is ordered by store, then numeric item ID.
    rows.sort(key=lambda x: (x[0], int(x[1])))
    return rows


def write_rfm(rows, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "RFM Data"
    ws.append(HEADERS)
    for store_id, item_id, in_price, del_price in rows:
        ws.append([store_id, str(item_id), in_price, del_price])  # Item ID stays text
    for i, h in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(len(h) + 2, 12)
    wb.save(out_path)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("input", type=Path)
    p.add_argument("-o", "--output", type=Path)
    p.add_argument("--sheet", help="Input sheet name (default: first sheet)")
    args = p.parse_args()

    out = args.output or args.input.with_name(args.input.stem + "_RFM_Format.xlsx")
    rows = read_wide(args.input, args.sheet)
    write_rfm(rows, out)
    print(f"Wrote {len(rows)} rows to {out}")
    # Each item ID once, in numeric order, joined by + with no spaces.
    print("+".join(sorted({r[1] for r in rows}, key=int)))


if __name__ == "__main__":
    main()
